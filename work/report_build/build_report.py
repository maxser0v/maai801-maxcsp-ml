from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
P=Path(__file__).resolve().parents[2]
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.1
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(17);styles['Heading 1'].paragraph_format.space_before=Pt(8)
styles['Heading 2'].font.size=Pt(12);styles['Heading 2'].paragraph_format.space_before=Pt(9)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.italic=True;styles['Caption'].font.bold=False
# Remove inherited decorative borders from every paragraph style.
for st in styles:
 for border in list(st.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='MAAI801     FOUNDATION OF AI     MINI PROJECT';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('MAAI801  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def clean(s):
 s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m: m.group(1) if m.group(1)==m.group(2) else m.group(1)+' ('+m.group(2)+')',s)
 return s.replace('`','').replace('**','').replace('*','')
def para(s,style=None):
 p=doc.add_paragraph(clean(s),style);return p
def table(lines):
 rows=[]
 for l in lines:
  cells=[clean(x.strip()) for x in l.strip().strip('|').split('|')]
  if all(re.fullmatch(r':?-+:?',x or '') for x in cells):continue
  rows.append(cells)
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(rows[0]);widths=[1.3,2.75,1.1,1.1] if n==4 else ([3.25,1.5,1.5] if n==3 else [6.25/n]*n)
 for j,c in enumerate(t.rows[0].cells):c.text=rows[0][j]
 for r in rows[1:]:
  cells=t.add_row().cells
  for j,text in enumerate(r):cells[j].text=text
 for ri,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17354A' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    m=OxmlElement('w:'+edge);m.set(qn('w:w'),'80');m.set(qn('w:type'),'dxa');margins.append(m)
   tcpr.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
    if j>0 and n==3:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
lines=(P/'work/report_draft.md').read_text().splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('>') or 'report draft' in line:continue
 if line.startswith('# '):
  para('Machine learning guided variable selection for Max CSP','Title')
  para('MAAI801 Foundation of AI     Experimental report','Subtitle');continue
 if line.startswith('## '):
  name=line[3:]
  if re.match(r'[2-7]\. ',name):doc.add_page_break()
  para(re.sub(r'[^\w\s]',' ',name),'Heading 1');continue
 if line.startswith('### '):para(re.sub(r'[^\w\s]',' ',re.sub(r'^[0-9.]+\s*','',line[4:])),'Heading 2');continue
 if line.startswith('|'):
  ls=[line]
  while i<len(lines) and lines[i].strip().startswith('|'):ls.append(lines[i]);i+=1
  table(ls);continue
 if line.startswith('**Figure'):
  match=re.search(r'figures/([^`]+\.png)',line)
  if match:
   fn=match.group(1)
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   p.add_run().add_picture(str(P/'figures'/fn),width=Inches(6.25))
   p.paragraph_format.keep_with_next=True
  caption=re.sub(r' Source image:.*','',line)
  para(caption,'Caption');continue
 if line.startswith('- '):para(line[2:],'List Bullet');continue
 para(line)
# Explicit scope disclosure without attributing work to nonexistent teammates.
para('Preparation note','Heading 2')
para('AI assistance generated the implementation, analysis, and draft deliverables. The submitting team must review the results, enter its actual member details and contribution record, and follow the course policy on AI assistance. No individual student contributions are asserted in this report.')
doc.core_properties.title='Machine learning guided variable selection for Max CSP'
doc.core_properties.author='MAAI801 project team'
doc.save(P/'submission'/'report.docx')
print(P/'submission'/'report.docx')
